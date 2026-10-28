# -*- coding: utf-8 -*-
"""词语积累 Markdown 同步为 Excel，不用 Word。"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path(__file__).resolve().parent
HEADERS = ["成语", "含义", "对象 / 侧重点", "易混", "相近成语"]
COL_WIDTHS = [14, 32, 30, 36, 22]
HEAD_FILL = PatternFill("solid", fgColor="1F6BB5")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=11)
IDIOM_FONT = Font(name="微软雅黑", size=11, bold=True, color="C0392B")
ALT_FILL = PatternFill("solid", fgColor="EBF5FB")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
WRAP = Alignment(wrap_text=True, vertical="center")


def strip_md(text: str) -> str:
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def is_four_chars(word: str) -> bool:
    return len(re.sub(r"\s+", "", word)) == 4


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.match(r"^:?-+:?$", c.replace(" ", "")) for c in cells):
            continue
        rows.append([strip_md(c) for c in cells])
    return rows


def collect_entries(md: str) -> tuple[list[dict], list[dict]]:
    four, two = [], []
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("|"):
            buf = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                buf.append(lines[i])
                i += 1
            parsed = parse_table(buf)
            if parsed:
                heads = parsed[0]
                for row in parsed[1:]:
                    mapped = {heads[j]: row[j] if j < len(row) else "" for j in range(len(heads))}
                    word = mapped.get("成语") or mapped.get("词") or (row[0] if row else "")
                    item = {
                        "成语": word,
                        "含义": mapped.get("含义", row[1] if len(row) > 1 else ""),
                        "对象 / 侧重点": mapped.get("对象 / 侧重点")
                        or mapped.get("搭配提示")
                        or mapped.get("和「视野」")
                        or (row[2] if len(row) > 2 else ""),
                        "易混": mapped.get("易混") or (row[3] if len(row) > 3 else ""),
                        "相近成语": mapped.get("相近成语") or (row[4] if len(row) > 4 else ""),
                    }
                    (four if is_four_chars(word) else two).append(item)
            continue
        i += 1
    return four, two


def write_sheet(ws, rows: list[dict], table_name: str):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    for i, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, i, h)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
        ws.column_dimensions[get_column_letter(i)].width = COL_WIDTHS[i - 1]
    ws.row_dimensions[1].height = 22
    last_col = get_column_letter(len(HEADERS))
    for r, item in enumerate(rows, 2):
        for c, key in enumerate(HEADERS, 1):
            cell = ws.cell(r, c, item.get(key, ""))
            cell.font = IDIOM_FONT if key == "成语" else BODY_FONT
            cell.alignment = WRAP
            cell.border = THIN
            if r % 2 == 0:
                cell.fill = ALT_FILL
        ws.row_dimensions[r].height = 36
    last = max(len(rows) + 1, 2)
    table = Table(displayName=table_name, ref=f"A1:{last_col}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def md_to_xlsx(md_path: Path, xlsx_path: Path):
    four, two = collect_entries(md_path.read_text(encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "四字成语"
    write_sheet(ws, four, "SiZiChengYu")
    if two:
        ws2 = wb.create_sheet("旧两字词")
        write_sheet(ws2, two, "JiuLiangZi")
    try:
        wb.save(str(xlsx_path))
        print("OK", xlsx_path.name)
    except PermissionError:
        alt = xlsx_path.with_name(xlsx_path.stem + "-新.xlsx")
        wb.save(str(alt))
        print("LOCKED", xlsx_path.name, "->", alt.name)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("-")]
    names = args or ["词语积累笔记"]
    for name in names:
        stem = name[:-3] if name.endswith(".md") else name
        md_path = BASE / f"{stem}.md"
        xlsx_path = BASE / f"{stem}.xlsx"
        if not md_path.exists():
            print("MISSING", md_path.name)
            continue
        md_to_xlsx(md_path, xlsx_path)


if __name__ == "__main__":
    main()
