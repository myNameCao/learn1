# -*- coding: utf-8 -*-
"""Markdown 笔记同步为 Word。"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from lxml import etree
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
COMPACT = False
H2_COLOR = RGBColor(0x1F, 0x6B, 0xB5)
H4_COLOR = RGBColor(0xC0, 0x39, 0x2B)
H5_COLOR = RGBColor(0x1A, 0x52, 0x76)
UP_COLOR = RGBColor(0x1E, 0x84, 0x4A)
DOWN_COLOR = RGBColor(0xC0, 0x39, 0x2B)
EQ_COLOR = RGBColor(0x1A, 0x52, 0x76)
HEAD_FILL = {
    "一、增长": "D4E6F1",
    "二、倍数": "D5F5E3",
    "三、比重": "FDEBD0",
    "四、平均数": "E8DAEF",
}
BODY_FILL = {
    "一、增长": "EBF5FB",
    "二、倍数": "EAF7F0",
    "三、比重": "FEF9F0",
    "四、平均数": "F5EEF8",
}
FORMULA_LABELS = (
    "平均数增长量",
    "平均数增长率",
    "总量增长率",
    "基期平均数",
    "比重增长量",
    "基期是几倍",
    "基期比重",
    "基期值",
    "现期值",
    "增长量",
    "增长率",
    "是几倍",
    "多几倍",
    "整体值",
    "部分值",
    "平均数",
    "总量",
    "份数",
    "比重",
)
HL_WORDS = (
    ("百分点", H4_COLOR, True),
    ("上升", UP_COLOR, True),
    ("增加", UP_COLOR, True),
    ("下降", DOWN_COLOR, True),
    ("减少", DOWN_COLOR, True),
    ("相等", EQ_COLOR, True),
    ("qA", H2_COLOR, True),
    ("qB", H2_COLOR, True),
)


def nest_indent_cm(level: int) -> float:
    if level >= 5:
        return 0.75
    if level >= 4:
        return 0.45
    return 0.0


def set_run_font(run, size=10.5, bold=False, color: RGBColor | None = None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_table_keep(table):
    for tr in table.rows:
        trPr = tr._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)


def append_omml_frac(paragraph, num: str, den: str, size=10.5):
    M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    def el(tag):
        return etree.Element(f"{{{M}}}{tag}")

    def math_run(text: str):
        r = el("r")
        rpr = el("rPr")
        sty = el("sty")
        sty.set(f"{{{M}}}val", "p")
        rpr.append(sty)
        r.append(rpr)
        wrpr = etree.Element(f"{{{W}}}rPr")
        sz = etree.Element(f"{{{W}}}sz")
        sz.set(f"{{{W}}}val", str(int(size * 2)))
        szcs = etree.Element(f"{{{W}}}szCs")
        szcs.set(f"{{{W}}}val", str(int(size * 2)))
        wrpr.append(sz)
        wrpr.append(szcs)
        r.append(wrpr)
        t = el("t")
        t.text = text
        r.append(t)
        return r

    omath = el("oMath")
    f = el("f")
    num_el = el("num")
    den_el = el("den")
    num_el.append(math_run(num))
    den_el.append(math_run(den))
    f.append(num_el)
    f.append(den_el)
    omath.append(f)
    paragraph._p.append(omath)


def append_hl_text(p, text, size=10.5, bold=False, color=None):
    if not text:
        return
    keys = [re.escape(w) for w, _, _ in HL_WORDS]
    parts = re.split(f"({'|'.join(keys)})", text) if keys else [text]
    cmap = {w: (c, b) for w, c, b in HL_WORDS}
    for part in parts:
        if not part:
            continue
        if part in cmap:
            c, b = cmap[part]
            r = p.add_run(part)
            set_run_font(r, size=size, bold=True if b else bold, color=c)
        else:
            r = p.add_run(part)
            set_run_font(r, size=size, bold=bold, color=color)


def append_text_with_frac(p, text, size=10.5, bold=False, color=None):
    pieces = re.split(r"(\\frac\{[^{}]+\}\{[^{}]+\})", text)
    for piece in pieces:
        if not piece:
            continue
        m = re.fullmatch(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", piece)
        if m:
            append_omml_frac(p, m.group(1), m.group(2), size=size)
        else:
            append_hl_text(p, piece, size=size, bold=bold, color=color)


def add_para(doc, text, size=10.5, bold=False, space_after=2, indent_cm=0.0, space_before=0, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    frac_after = 1 if COMPACT else 8
    frac_ls = 1.12 if COMPACT else 1.6
    body_after = 0 if COMPACT else space_after
    p.paragraph_format.space_after = Pt(max(body_after, frac_after) if "\\frac" in text else body_after)
    if "\\frac" in text:
        p.paragraph_format.line_spacing = frac_ls
    else:
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if indent_cm:
        p.paragraph_format.left_indent = Cm(indent_cm)
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            set_run_font(r, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            set_run_font(r, size=max(size - 1, 8))
        else:
            append_text_with_frac(p, part, size=size, bold=bold, color=color)
    return p


def set_table_indent(table, indent_cm: float):
    if not indent_cm:
        return
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblInd = OxmlElement("w:tblInd")
    tblInd.set(qn("w:w"), str(int(indent_cm * 567)))
    tblInd.set(qn("w:type"), "dxa")
    tblPr.append(tblInd)


def fill_rich_cell(cell, txt, font_size, header=False):
    chunks = [c.strip() for c in re.split(r"<br\s*/?>", txt, flags=re.I)]
    if not chunks:
        chunks = [""]
    for i, chunk in enumerate(chunks):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1 if COMPACT else 2)
        p.paragraph_format.line_spacing = 1.3 if "\\frac" in chunk else 1.0
        labeled = False
        if not header:
            for lab in FORMULA_LABELS:
                if chunk.startswith(lab + " ") or chunk.startswith(lab + "　") or chunk.startswith(lab + "="):
                    r = p.add_run(lab)
                    set_run_font(r, size=font_size, bold=True, color=H4_COLOR)
                    chunk = chunk[len(lab) :]
                    labeled = True
                    break
        parts = re.split(r"(\*\*[^*]+\*\*)", chunk)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                inner = part[2:-2]
                r = p.add_run(inner)
                if inner in (">",):
                    set_run_font(r, size=font_size, bold=True, color=UP_COLOR)
                elif inner in ("<",):
                    set_run_font(r, size=font_size, bold=True, color=DOWN_COLOR)
                elif inner in ("=",):
                    set_run_font(r, size=font_size, bold=True, color=EQ_COLOR)
                elif "求" in inner or "判断" in inner or "基本" in inner or inner.startswith("（"):
                    set_run_font(r, size=font_size, bold=True, color=H4_COLOR)
                else:
                    set_run_font(r, size=font_size, bold=True, color=H4_COLOR)
            else:
                append_text_with_frac(p, part, size=font_size, bold=header)


def flush_table(doc, rows, indent_cm=0.0):
    parsed = []
    for row in rows:
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        if all(re.match(r"^:?-+:?$", c.replace(" ", "")) for c in cells):
            continue
        parsed.append(cells)
    if not parsed:
        return
    cols = max(len(r) for r in parsed)
    t = doc.add_table(rows=len(parsed), cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    usable = Cm(17.4 if COMPACT else 17.0)
    col_w = usable / cols
    if COMPACT and cols == 2:
        font_size = 8.5
    elif cols >= 4:
        font_size = 8
    else:
        font_size = 9
    for i, row in enumerate(parsed):
        header = i == 0 or (
            COMPACT and cols == 2 and all(len(c) <= 8 and "\\" not in c for c in row)
        )
        for j in range(cols):
            cell = t.cell(i, j)
            cell.width = col_w
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            txt = row[j] if j < len(row) else ""
            fill_rich_cell(cell, txt, font_size, header=header)
            col_name = parsed[0][j] if parsed and j < len(parsed[0]) else ""
            if header:
                set_cell_shading(cell, HEAD_FILL.get(col_name, "F0F0F0"))
            elif col_name in BODY_FILL:
                set_cell_shading(cell, BODY_FILL[col_name])
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcMar = OxmlElement("w:tcMar")
            pad = "40" if COMPACT else "20"
            for side, val in (("top", pad), ("bottom", pad), ("left", "50"), ("right", "50")):
                node = OxmlElement(f"w:{side}")
                node.set(qn("w:w"), val)
                node.set(qn("w:type"), "dxa")
                tcMar.append(node)
            tcPr.append(tcMar)
    set_table_keep(t)
    set_table_indent(t, indent_cm)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    if indent_cm:
        sp.paragraph_format.left_indent = Cm(indent_cm)


def md_to_docx(md_path: Path, docx_path: Path):
    global COMPACT
    COMPACT = md_path.stem == "常用公式"
    md = md_path.read_text(encoding="utf-8")
    doc = Document()
    for section in doc.sections:
        m = 0.55 if COMPACT else 1.0
        side = 0.8 if COMPACT else 1.2
        section.top_margin = Cm(m)
        section.bottom_margin = Cm(m)
        section.left_margin = Cm(side)
        section.right_margin = Cm(side)
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(10.5)

    in_code = False
    code_buf = []
    in_formula = False
    formula_buf = []
    heading_level = 1
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("$$"):
            if not in_formula:
                in_formula = True
                formula_buf = []
                rest = line.strip()[2:].strip()
                if rest.endswith("$$"):
                    add_para(doc, rest[:-2].strip(), size=10)
                    in_formula = False
                elif rest:
                    formula_buf.append(rest)
            else:
                add_para(doc, " ".join(formula_buf), size=10)
                in_formula = False
                formula_buf = []
            i += 1
            continue
        if in_formula:
            if line.strip().endswith("$$"):
                formula_buf.append(line.strip()[:-2])
                add_para(doc, " ".join(formula_buf), size=10)
                in_formula = False
                formula_buf = []
            else:
                formula_buf.append(line.strip())
            i += 1
            continue
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                p = doc.add_paragraph()
                r = p.add_run("\n".join(code_buf))
                set_run_font(r, size=9)
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if line.strip().startswith("|"):
            table_buf = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_buf.append(lines[i])
                i += 1
            flush_table(doc, table_buf, indent_cm=nest_indent_cm(heading_level))
            continue
        if line.startswith("# "):
            heading_level = 1
            title = line[2:].strip()
            if COMPACT:
                is_cover = title == "常用公式"
                add_para(
                    doc,
                    title,
                    size=13 if is_cover else 11,
                    bold=True,
                    space_before=0 if is_cover else 5,
                    space_after=1,
                    color=None if is_cover else H2_COLOR,
                )
            else:
                p = doc.add_heading(title, level=1)
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(4)
                for r in p.runs:
                    set_run_font(r, size=16, bold=True)
        elif line.startswith("## "):
            heading_level = 2
            title = line[3:].strip()
            if COMPACT:
                add_para(
                    doc,
                    title,
                    size=10,
                    bold=True,
                    space_before=3,
                    space_after=0,
                    color=H4_COLOR,
                )
            else:
                p = doc.add_heading(title, level=2)
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                for r in p.runs:
                    set_run_font(r, size=13, bold=True, color=H2_COLOR)
        elif line.startswith("### "):
            heading_level = 3
            p = doc.add_heading(line[4:].strip(), level=3)
            p.paragraph_format.space_before = Pt(4 if COMPACT else 8)
            p.paragraph_format.space_after = Pt(2 if COMPACT else 3)
            for r in p.runs:
                set_run_font(r, size=11 if COMPACT else 12, bold=True)
        elif line.startswith("#### "):
            heading_level = 4
            p = add_para(
                doc,
                line[5:].strip(),
                size=12,
                bold=True,
                space_after=3,
                space_before=8,
                indent_cm=0.2,
                color=H4_COLOR,
            )
        elif line.startswith("##### "):
            heading_level = 5
            p = add_para(
                doc,
                line[6:].strip(),
                size=11,
                bold=True,
                space_after=2,
                space_before=6,
                indent_cm=0.45,
                color=H5_COLOR,
            )
        elif line.strip() in ("", "---"):
            pass
        elif line.strip().startswith(">"):
            add_para(
                doc,
                line.strip().lstrip("> ").strip(),
                size=9.5,
                space_after=2,
                indent_cm=nest_indent_cm(heading_level),
            )
        elif (
            line.strip().startswith("*")
            and line.strip().endswith("*")
            and not line.strip().startswith("**")
        ):
            add_para(doc, line.strip()[1:-1], size=9.5, indent_cm=nest_indent_cm(heading_level))
        else:
            add_para(
                doc,
                line,
                size=9 if COMPACT else 10.5,
                indent_cm=nest_indent_cm(heading_level),
            )
        i += 1

    try:
        doc.save(str(docx_path))
        print("OK", docx_path.name)
    except PermissionError:
        alt = docx_path.with_name(docx_path.stem + "-新.docx")
        doc.save(str(alt))
        print("LOCKED", docx_path.name, "->", alt.name)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("-")]
    if args:
        names = args
    else:
        names = ["言语理解笔记"]
    for name in names:
        stem = name[:-3] if name.endswith(".md") else name
        md_path = BASE / f"{stem}.md"
        docx_path = BASE / f"{stem}.docx"
        if stem == "词语积累笔记":
            print("SKIP", md_path.name, "(use 导出Excel.py)")
            continue
        if not md_path.exists():
            print("MISSING", md_path.name)
            continue
        md_to_docx(md_path, docx_path)


if __name__ == "__main__":
    main()
